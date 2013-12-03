# encoding: utf-8

"""Test suite for pptx.text module."""

from __future__ import absolute_import, print_function

import pytest

from pptx.constants import MSO, PP
from pptx.dml.color import RGBColor
from pptx.enum import MSO_COLOR_TYPE, MSO_THEME_COLOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml import parse_xml_bytes
from pptx.oxml.ns import namespaces, nsdecls
from pptx.oxml.text import (
    CT_RegularTextRun, CT_TextCharacterProperties, CT_TextParagraph
)
from pptx.text import (
    _Font, _FontColor, _Hyperlink, _Paragraph, _Run, TextFrame
)
from pptx.util import Inches

from .oxml.unitdata.dml import (
    a_lumMod, a_lumOff, a_schemeClr, a_solidFill, an_srgbClr
)
from .oxml.unitdata.text import (
    a_bodyPr, a_txBody, a_p, a_pPr, a_t, an_hlinkClick, an_r, an_rPr
)
from .unitutil import (
    absjoin, actual_xml, class_mock, instance_mock, loose_mock,
    parse_xml_file, property_mock, test_file_dir
)


nsmap = namespaces('a', 'r', 'p')


class DescribeTextFrame(object):

    def it_knows_the_number_of_paragraphs_it_contains(
            self, txBody, txBody_with_2_paras):
        assert len(TextFrame(txBody, None).paragraphs) == 1
        assert len(TextFrame(txBody_with_2_paras, None).paragraphs) == 2

    def it_can_add_a_paragraph_to_the_text_it_contains(
            self, txBody, txBody_with_2_paras_xml):
        textframe = TextFrame(txBody, None)
        textframe.add_paragraph()
        assert actual_xml(textframe._txBody) == txBody_with_2_paras_xml

    def it_can_replace_the_text_it_contains(
            self, txBody, txBody_with_text_xml):
        textframe = TextFrame(txBody, None)
        textframe.text = 'foobar'
        assert actual_xml(txBody) == txBody_with_text_xml

    def it_can_get_its_margin_settings(
            self, txBody, txBody_with_lIns, txBody_with_tIns,
            txBody_with_rIns, txBody_with_bIns):

        textframe = TextFrame(txBody, None)
        assert textframe.margin_left is None
        assert textframe.margin_top is None
        assert textframe.margin_right is None
        assert textframe.margin_bottom is None

        textframe = TextFrame(txBody_with_lIns, None)
        assert textframe.margin_left == Inches(0.01)

        textframe = TextFrame(txBody_with_tIns, None)
        assert textframe.margin_top == Inches(0.02)

        textframe = TextFrame(txBody_with_rIns, None)
        assert textframe.margin_right == Inches(0.03)

        textframe = TextFrame(txBody_with_bIns, None)
        assert textframe.margin_bottom == Inches(0.04)

    def it_can_change_its_margin_settings(
            self, txBody_bldr, txBody_with_lIns_xml, txBody_with_tIns_xml,
            txBody_with_rIns_xml, txBody_with_bIns_xml):

        textframe = TextFrame(txBody_bldr.element, None)
        textframe.margin_left = Inches(0.01)
        assert actual_xml(textframe._txBody) == txBody_with_lIns_xml

        textframe = TextFrame(txBody_bldr.element, None)
        textframe.margin_top = Inches(0.02)
        assert actual_xml(textframe._txBody) == txBody_with_tIns_xml

        textframe = TextFrame(txBody_bldr.element, None)
        textframe.margin_right = Inches(0.03)
        assert actual_xml(textframe._txBody) == txBody_with_rIns_xml

        textframe = TextFrame(txBody_bldr.element, None)
        textframe.margin_bottom = Inches(0.04)
        assert actual_xml(textframe._txBody) == txBody_with_bIns_xml

    def it_raises_on_attempt_to_set_margin_to_non_int(self, textframe):
        with pytest.raises(ValueError):
            textframe.margin_bottom = '0.1'

    def it_can_change_its_vertical_anchor_setting(
            self, txBody, txBody_with_anchor_ctr_xml):
        textframe = TextFrame(txBody, None)
        textframe.vertical_anchor = MSO.ANCHOR_MIDDLE
        assert actual_xml(textframe._txBody) == txBody_with_anchor_ctr_xml

    def it_can_change_the_word_wrap_setting(
            self, txBody, txBody_with_wrap_on_xml, txBody_with_wrap_off_xml,
            txBody_xml):
        textframe = TextFrame(txBody, None)
        assert textframe.word_wrap is None

        textframe.word_wrap = True
        assert actual_xml(textframe._txBody) == txBody_with_wrap_on_xml
        assert textframe.word_wrap is True

        textframe.word_wrap = False
        assert actual_xml(textframe._txBody) == txBody_with_wrap_off_xml
        assert textframe.word_wrap is False

        textframe.word_wrap = None
        assert actual_xml(textframe._txBody) == txBody_xml
        assert textframe.word_wrap is None

    def it_knows_the_part_it_belongs_to(self, textframe_with_parent_):
        textframe, parent_ = textframe_with_parent_
        part = textframe.part
        assert part is parent_.part

    # fixtures ---------------------------------------------

    @pytest.fixture
    def textframe(self, txBody):
        return TextFrame(txBody, None)

    @pytest.fixture
    def textframe_with_parent_(self, request):
        parent_ = loose_mock(request, name='parent_')
        textframe = TextFrame(None, parent_)
        return textframe, parent_

    @pytest.fixture
    def txBody(self, txBody_bldr):
        return txBody_bldr.element

    @pytest.fixture
    def txBody_bldr(self):
        p_bldr = a_p()
        bodyPr_bldr = a_bodyPr()
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
        )

    @pytest.fixture
    def txBody_xml(self, txBody_bldr):
        return txBody_bldr.xml()

    @pytest.fixture
    def txBody_with_2_paras(self, _txBody_with_2_paras_bldr):
        return _txBody_with_2_paras_bldr.element

    @pytest.fixture
    def txBody_with_2_paras_xml(self, _txBody_with_2_paras_bldr):
        return _txBody_with_2_paras_bldr.xml()

    @pytest.fixture
    def txBody_with_anchor_ctr_xml(self):
        p_bldr = a_p()
        bodyPr_bldr = a_bodyPr().with_anchor('ctr')
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
                      .xml()
        )

    @pytest.fixture
    def txBody_with_bIns(self, _txBody_with_bIns_bldr):
        return _txBody_with_bIns_bldr.element

    @pytest.fixture
    def txBody_with_bIns_xml(self, _txBody_with_bIns_bldr):
        return _txBody_with_bIns_bldr.xml()

    @pytest.fixture
    def txBody_with_lIns(self, _txBody_with_lIns_bldr):
        return _txBody_with_lIns_bldr.element

    @pytest.fixture
    def txBody_with_lIns_xml(self, _txBody_with_lIns_bldr):
        return _txBody_with_lIns_bldr.xml()

    @pytest.fixture
    def txBody_with_rIns(self, _txBody_with_rIns_bldr):
        return _txBody_with_rIns_bldr.element

    @pytest.fixture
    def txBody_with_rIns_xml(self, _txBody_with_rIns_bldr):
        return _txBody_with_rIns_bldr.xml()

    @pytest.fixture
    def txBody_with_tIns(self, _txBody_with_tIns_bldr):
        return _txBody_with_tIns_bldr.element

    @pytest.fixture
    def txBody_with_tIns_xml(self, _txBody_with_tIns_bldr):
        return _txBody_with_tIns_bldr.xml()

    @pytest.fixture
    def txBody_with_text_xml(self):
        t_bldr = a_t().with_text('foobar')
        r_bldr = an_r().with_child(t_bldr)
        p_bldr = a_p().with_child(r_bldr)
        bodyPr_bldr = a_bodyPr()
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
                      .xml()
        )

    @pytest.fixture
    def txBody_with_wrap_off_xml(self):
        p_bldr = a_p()
        bodyPr_bldr = a_bodyPr().with_wrap('none')
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
                      .xml()
        )

    @pytest.fixture
    def txBody_with_wrap_on_xml(self):
        p_bldr = a_p()
        bodyPr_bldr = a_bodyPr().with_wrap('square')
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
                      .xml()
        )

    @pytest.fixture
    def _txBody_with_2_paras_bldr(self):
        p_bldr = a_p()
        bodyPr_bldr = a_bodyPr()
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
                      .with_child(p_bldr)
        )

    @pytest.fixture
    def _txBody_with_bIns_bldr(self):
        p_bldr = a_p()
        bodyPr_bldr = a_bodyPr().with_bIns(int(914400*0.04))
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
        )

    @pytest.fixture
    def _txBody_with_lIns_bldr(self):
        p_bldr = a_p()
        bodyPr_bldr = a_bodyPr().with_lIns(int(914400*0.01))
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
        )

    @pytest.fixture
    def _txBody_with_tIns_bldr(self):
        p_bldr = a_p()
        bodyPr_bldr = a_bodyPr().with_tIns(int(914400*0.02))
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
        )

    @pytest.fixture
    def _txBody_with_rIns_bldr(self):
        p_bldr = a_p()
        bodyPr_bldr = a_bodyPr().with_rIns(int(914400*0.03))
        return (
            a_txBody().with_nsdecls()
                      .with_child(bodyPr_bldr)
                      .with_child(p_bldr)
        )


class Describe_Font(object):

    def it_knows_the_bold_setting(self, font, bold_font, bold_off_font):
        assert font.bold is None
        assert bold_font.bold is True
        assert bold_off_font.bold is False

    def it_can_change_the_bold_setting(
            self, font, bold_rPr_xml, bold_off_rPr_xml, rPr_xml):
        assert actual_xml(font._rPr) == rPr_xml
        font.bold = None
        assert actual_xml(font._rPr) == rPr_xml
        font.bold = True
        assert actual_xml(font._rPr) == bold_rPr_xml
        font.bold = False
        assert actual_xml(font._rPr) == bold_off_rPr_xml
        font.bold = None
        assert actual_xml(font._rPr) == rPr_xml

    def it_has_a_color(self, font):
        assert isinstance(font.color, _FontColor)

    def it_knows_the_italic_setting(self, font, italic_font, italic_off_font):
        assert font.italic is None
        assert italic_font.italic is True
        assert italic_off_font.italic is False

    def it_can_change_the_italic_setting(
            self, font, italic_rPr_xml, italic_off_rPr_xml, rPr_xml):
        assert actual_xml(font._rPr) == rPr_xml
        font.italic = None  # important to test None to None transition
        assert actual_xml(font._rPr) == rPr_xml
        font.italic = True
        assert actual_xml(font._rPr) == italic_rPr_xml
        font.italic = False
        assert actual_xml(font._rPr) == italic_off_rPr_xml
        font.italic = None
        assert actual_xml(font._rPr) == rPr_xml

    def it_can_set_the_font_size(self, font):
        font.size = 2400
        expected_xml = an_rPr().with_nsdecls().with_sz(2400).xml()
        assert actual_xml(font._rPr) == expected_xml

    # fixtures ---------------------------------------------

    @pytest.fixture
    def bold_font(self):
        bold_rPr = an_rPr().with_nsdecls().with_b(1).element
        return _Font(bold_rPr)

    @pytest.fixture
    def bold_off_font(self):
        bold_off_rPr = an_rPr().with_nsdecls().with_b(0).element
        return _Font(bold_off_rPr)

    @pytest.fixture
    def bold_off_rPr_xml(self):
        return an_rPr().with_nsdecls().with_b(0).xml()

    @pytest.fixture
    def bold_rPr_xml(self):
        return an_rPr().with_nsdecls().with_b(1).xml()

    @pytest.fixture
    def italic_font(self):
        italic_rPr = an_rPr().with_nsdecls().with_i(1).element
        return _Font(italic_rPr)

    @pytest.fixture
    def italic_off_font(self):
        italic_rPr = an_rPr().with_nsdecls().with_i(0).element
        return _Font(italic_rPr)

    @pytest.fixture
    def italic_off_rPr_xml(self):
        return an_rPr().with_nsdecls().with_i(0).xml()

    @pytest.fixture
    def italic_rPr_xml(self):
        return an_rPr().with_nsdecls().with_i(1).xml()

    @pytest.fixture
    def rPr_xml(self):
        return an_rPr().with_nsdecls().xml()

    @pytest.fixture
    def font(self):
        rPr = an_rPr().with_nsdecls().element
        return _Font(rPr)


class Describe_FontColor(object):

    def it_knows_the_type_of_its_color(
            self, color, rgb_color, theme_color, solidFill_only_color):
        assert color.type is None
        assert rgb_color.type is MSO_COLOR_TYPE.RGB
        assert theme_color.type is MSO_COLOR_TYPE.SCHEME
        assert solidFill_only_color.type is None

    def it_knows_the_RGB_value_of_an_RGB_color(self, color, rgb_color):
        assert color.rgb is None
        assert rgb_color.rgb == RGBColor(0x12, 0x34, 0x56)

    def it_knows_the_theme_color_of_a_theme_color(self, color, theme_color):
        assert color.theme_color is None
        assert theme_color.theme_color == MSO_THEME_COLOR.ACCENT_1

    def it_knows_the_brightness_adjustment_of_its_color(
            self, color, solidFill_only_color, rgb_color, theme_color,
            rgb_color_with_brightness, theme_color_with_brightness):
        assert color.brightness == 0.0
        assert solidFill_only_color.brightness == 0.0
        assert rgb_color.brightness == 0.0
        assert theme_color.brightness == 0.0
        assert rgb_color_with_brightness.brightness == 0.4
        assert theme_color_with_brightness.brightness == -0.25

    def it_can_set_the_RGB_color_value(self, color, rgb_color, theme_color):
        color.rgb = RGBColor(0x01, 0x23, 0x45)
        assert color.rgb == RGBColor(0x01, 0x23, 0x45)
        rgb_color.rgb = RGBColor(0x67, 0x89, 0x9A)
        assert rgb_color.rgb == RGBColor(0x67, 0x89, 0x9A)
        theme_color.rgb = RGBColor(0xBC, 0xDE, 0xF0)
        assert theme_color.rgb == RGBColor(0xBC, 0xDE, 0xF0)

    def it_can_set_the_theme_color(self, color, rgb_color, theme_color):
        color.theme_color = MSO_THEME_COLOR.ACCENT_1
        assert color.theme_color == MSO_THEME_COLOR.ACCENT_1
        rgb_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
        assert rgb_color.theme_color == MSO_THEME_COLOR.BACKGROUND_1
        theme_color.theme_color = MSO_THEME_COLOR.TEXT_1
        assert theme_color.theme_color == MSO_THEME_COLOR.TEXT_1

    def it_can_set_the_color_brightness(self, color, rgb_color, theme_color):
        rgb_color.brightness = 0.4
        assert rgb_color.brightness == 0.4
        theme_color.brightness = 0.0
        assert theme_color.brightness == 0.0
        theme_color.brightness = -0.25
        assert theme_color.brightness == -0.25

    def it_raises_on_attempt_to_set_brightness_out_of_range(
            self, theme_color):
        with pytest.raises(ValueError):
            theme_color.brightness = 1.1
        with pytest.raises(ValueError):
            theme_color.brightness = -1.1
        with pytest.raises(ValueError):
            theme_color.brightness = 'bright'

    def it_raises_on_attempt_to_set_brightness_on_None_color_type(
            self, color):
        with pytest.raises(ValueError):
            color.brightness = 0

    # fixtures ---------------------------------------------

    @pytest.fixture
    def color(self):
        rPr = an_rPr().with_nsdecls().element
        return _FontColor(rPr)

    @pytest.fixture
    def rgb_color(self):
        srgbClr_bldr = an_srgbClr().with_val('123456')
        solidFill_bldr = a_solidFill().with_child(srgbClr_bldr)
        rPr = an_rPr().with_nsdecls().with_child(solidFill_bldr).element
        return _FontColor(rPr)

    @pytest.fixture
    def rgb_color_with_brightness(self):
        lumMod_bldr = a_lumMod().with_val('60000')
        lumOff_bldr = a_lumOff().with_val('40000')
        srgbClr_bldr = an_srgbClr().with_val('123456')
        srgbClr_bldr.with_child(lumMod_bldr).with_child(lumOff_bldr)
        solidFill_bldr = a_solidFill().with_child(srgbClr_bldr)
        rPr = an_rPr().with_nsdecls().with_child(solidFill_bldr).element
        return _FontColor(rPr)

    @pytest.fixture
    def solidFill_only_color(self):
        solidFill_bldr = a_solidFill()
        rPr = an_rPr().with_nsdecls().with_child(solidFill_bldr).element
        return _FontColor(rPr)

    @pytest.fixture
    def theme_color(self):
        schemeClr_bldr = a_schemeClr().with_val('accent1')
        solidFill_bldr = a_solidFill().with_child(schemeClr_bldr)
        rPr = an_rPr().with_nsdecls().with_child(solidFill_bldr).element
        return _FontColor(rPr)

    @pytest.fixture
    def theme_color_with_brightness(self):
        lumMod_bldr = a_lumMod().with_val('75000')
        schemeClr_bldr = a_schemeClr().with_val('foobar')
        schemeClr_bldr.with_child(lumMod_bldr)
        solidFill_bldr = a_solidFill().with_child(schemeClr_bldr)
        rPr = an_rPr().with_nsdecls().with_child(solidFill_bldr).element
        return _FontColor(rPr)


class Describe_Hyperlink(object):

    def it_knows_the_target_url_of_the_hyperlink(self, hlink_with_url_):
        hlink, rId, url = hlink_with_url_
        assert hlink.address == url
        hlink.part.target_ref.assert_called_once_with(rId)

    def it_has_None_for_address_when_no_hyperlink_is_present(self, hlink):
        assert hlink.address is None

    def it_can_set_the_target_url(
            self, hlink, rPr_with_hlinkClick_xml, url):
        hlink.address = url
        # verify -----------------------
        hlink.part.relate_to.assert_called_once_with(
            url, RT.HYPERLINK, is_external=True
        )
        assert actual_xml(hlink._rPr) == rPr_with_hlinkClick_xml
        assert hlink.address == url

    def it_can_remove_the_hyperlink(self, remove_hlink_fixture_):
        hlink, rPr_xml, rId = remove_hlink_fixture_
        hlink.address = None
        assert actual_xml(hlink._rPr) == rPr_xml
        hlink.part.drop_rel.assert_called_once_with(rId)

    def it_should_remove_the_hyperlink_when_url_set_to_empty_string(
            self, remove_hlink_fixture_):
        hlink, rPr_xml, rId = remove_hlink_fixture_
        hlink.address = ''
        assert actual_xml(hlink._rPr) == rPr_xml
        hlink.part.drop_rel.assert_called_once_with(rId)

    def it_can_change_the_target_url(self, change_hlink_fixture_):
        # fixture ----------------------
        hlink, rId_existing, new_url, new_rPr_xml = change_hlink_fixture_
        # exercise ---------------------
        hlink.address = new_url
        # verify -----------------------
        assert actual_xml(hlink._rPr) == new_rPr_xml
        hlink.part.drop_rel.assert_called_once_with(rId_existing)
        hlink.part.relate_to.assert_called_once_with(
            new_url, RT.HYPERLINK, is_external=True
        )

    # fixtures ---------------------------------------------

    @pytest.fixture
    def change_hlink_fixture_(
            self, request, hlink_with_hlinkClick, rId, rId_2, part_, url_2):
        hlinkClick_bldr = an_hlinkClick().with_rId(rId_2)
        new_rPr_xml = (
            an_rPr().with_nsdecls('a', 'r')
                    .with_child(hlinkClick_bldr)
                    .xml()
        )
        part_.relate_to.return_value = rId_2
        property_mock(request, _Hyperlink, 'part', return_value=part_)
        return hlink_with_hlinkClick, rId, url_2, new_rPr_xml

    @pytest.fixture
    def hlink(self, request, part_):
        rPr = an_rPr().with_nsdecls('a', 'r').element
        hlink = _Hyperlink(rPr, None)
        property_mock(request, _Hyperlink, 'part', return_value=part_)
        return hlink

    @pytest.fixture
    def hlink_with_hlinkClick(self, request, rPr_with_hlinkClick_bldr):
        rPr = rPr_with_hlinkClick_bldr.element
        return _Hyperlink(rPr, None)

    @pytest.fixture
    def hlink_with_url_(
            self, request, part_, hlink_with_hlinkClick, rId, url):
        property_mock(request, _Hyperlink, 'part', return_value=part_)
        return hlink_with_hlinkClick, rId, url

    @pytest.fixture
    def part_(self, request, url, rId):
        """
        Mock Part instance suitable for patching into _Hyperlink.part
        property. It returns url for target_ref() and rId for relate_to().
        """
        part_ = instance_mock(request, Part)
        part_.target_ref.return_value = url
        part_.relate_to.return_value = rId
        return part_

    @pytest.fixture
    def rId(self):
        return 'rId2'

    @pytest.fixture
    def rId_2(self):
        return 'rId6'

    @pytest.fixture
    def remove_hlink_fixture_(
            self, request, hlink_with_hlinkClick, rPr_xml, rId):
        property_mock(request, _Hyperlink, 'part')
        return hlink_with_hlinkClick, rPr_xml, rId

    @pytest.fixture
    def rPr_with_hlinkClick_bldr(self, rId):
        hlinkClick_bldr = an_hlinkClick().with_rId(rId)
        rPr_bldr = (
            an_rPr().with_nsdecls('a', 'r')
                    .with_child(hlinkClick_bldr)
        )
        return rPr_bldr

    @pytest.fixture
    def rPr_with_hlinkClick_xml(self, rPr_with_hlinkClick_bldr):
        return rPr_with_hlinkClick_bldr.xml()

    @pytest.fixture
    def rPr_xml(self):
        return an_rPr().with_nsdecls().xml()

    @pytest.fixture
    def url(self):
        return 'https://github.com/scanny/python-pptx'

    @pytest.fixture
    def url_2(self):
        return 'https://pypi.python.org/pypi/python-pptx'


class Describe_Paragraph(object):

    def it_can_add_a_run(self, paragraph, p_with_r_xml):
        run = paragraph.add_run()
        assert actual_xml(paragraph._p) == p_with_r_xml
        assert isinstance(run, _Run)

    def it_knows_the_alignment_setting_of_the_paragraph(
            self, paragraph, paragraph_with_algn):
        assert paragraph.alignment is None
        assert paragraph_with_algn.alignment == PP.ALIGN_CENTER

    def it_can_change_its_alignment_setting(self, paragraph):
        paragraph.alignment = PP.ALIGN_LEFT
        assert paragraph._pPr.algn == 'l'
        paragraph.alignment = None
        assert paragraph._pPr.algn is None

    def it_can_delete_the_text_it_contains(self, paragraph, p_):
        paragraph._p = p_
        paragraph.clear()
        p_.remove_child_r_elms.assert_called_once_with()

    def it_provides_access_to_the_default_paragraph_font(
            self, paragraph, Font_):
        font = paragraph.font
        Font_.assert_called_once_with(paragraph._defRPr)
        assert font == Font_.return_value

    def test_level_setter_generates_correct_xml(self, paragraph_with_text):
        """_Paragraph.level setter generates correct XML"""
        # setup ------------------------
        expected_xml = (
            '<a:p %s>\n  <a:pPr lvl="2"/>\n  <a:r>\n    <a:t>test text</a:t>'
            '\n  </a:r>\n</a:p>\n' % nsdecls('a')
        )
        # exercise ---------------------
        paragraph_with_text.level = 2
        # verify -----------------------
        assert actual_xml(paragraph_with_text._p) == expected_xml

    def test_level_default_is_zero(self, paragraph_with_text):
        """_Paragraph.level defaults to zero on no lvl attribute"""
        # verify -----------------------
        assert paragraph_with_text.level == 0

    def test_level_roundtrips_intact(self, paragraph_with_text):
        """_Paragraph.level property round-trips intact"""
        # exercise ---------------------
        paragraph_with_text.level = 5
        # verify -----------------------
        assert paragraph_with_text.level == 5

    def test_level_raises_on_bad_value(self, paragraph_with_text):
        """_Paragraph.level raises on attempt to assign invalid value"""
        test_cases = ('0', -1, 9)
        for value in test_cases:
            with pytest.raises(ValueError):
                paragraph_with_text.level = value

    def test_runs_size(self, pList):
        """_Paragraph.runs is expected size"""
        # setup ------------------------
        actual_lengths = []
        for p in pList:
            paragraph = _Paragraph(p, None)
            # exercise ----------------
            actual_lengths.append(len(paragraph.runs))
        # verify ------------------
        expected = [0, 0, 2, 1, 1, 1]
        actual = actual_lengths
        msg = "expected run count %s, got %s" % (expected, actual)
        assert actual == expected, msg

    def test_text_setter_sets_single_run_text(self, pList):
        """assignment to _Paragraph.text creates single run containing value"""
        # setup ------------------------
        test_text = 'python-pptx was here!!'
        p_elm = pList[2]
        paragraph = _Paragraph(p_elm, None)
        # exercise ---------------------
        paragraph.text = test_text
        # verify -----------------------
        assert len(paragraph.runs) == 1
        assert paragraph.runs[0].text == test_text

    def test_text_accepts_non_ascii_strings(self, paragraph_with_text):
        """assignment of non-ASCII string to text does not raise"""
        # setup ------------------------
        _7bit_string = 'String containing only 7-bit (ASCII) characters'
        _8bit_string = '8-bit string: Hér er texti með íslenskum stöfum.'
        _utf8_literal = u'unicode literal: Hér er texti með íslenskum stöfum.'
        _utf8_from_8bit = unicode('utf-8 unicode: Hér er texti', 'utf-8')
        # verify -----------------------
        try:
            text = _7bit_string
            paragraph_with_text.text = text
            text = _8bit_string
            paragraph_with_text.text = text
            text = _utf8_literal
            paragraph_with_text.text = text
            text = _utf8_from_8bit
            paragraph_with_text.text = text
        except ValueError:
            msg = "_Paragraph.text rejects valid text string '%s'" % text
            pytest.fail(msg)

    # fixtures ---------------------------------------------

    @pytest.fixture
    def Font_(self, request):
        return class_mock(request, 'pptx.text._Font')

    @pytest.fixture
    def pList(self, sld, xpath):
        return sld.xpath(xpath, namespaces=nsmap)

    @pytest.fixture
    def p_(self, request):
        return instance_mock(request, CT_TextParagraph)

    @pytest.fixture
    def p_bldr(self):
        return a_p().with_nsdecls()

    @pytest.fixture
    def p_with_r_xml(self):
        run_bldr = an_r().with_child(a_t())
        return a_p().with_nsdecls().with_child(run_bldr).xml()

    @pytest.fixture
    def p_with_text(self, p_with_text_xml):
        return parse_xml_bytes(p_with_text_xml)

    @pytest.fixture
    def p_with_text_xml(self, test_text):
        return ('<a:p %s><a:r><a:t>%s</a:t></a:r></a:p>' %
                (nsdecls('a'), test_text))

    @pytest.fixture
    def paragraph(self, p_bldr):
        return _Paragraph(p_bldr.element, None)

    @pytest.fixture
    def paragraph_with_algn(self):
        pPr_bldr = a_pPr().with_algn('ctr')
        p_bldr = a_p().with_nsdecls().with_child(pPr_bldr)
        return _Paragraph(p_bldr.element, None)

    @pytest.fixture
    def paragraph_with_text(self, p_with_text):
        return _Paragraph(p_with_text, None)

    @pytest.fixture
    def path(self):
        return absjoin(test_file_dir, 'slide1.xml')

    @pytest.fixture
    def sld(self, path):
        return parse_xml_file(path).getroot()

    @pytest.fixture
    def test_text(self):
        return 'test text'

    @pytest.fixture
    def xpath(self):
        return './p:cSld/p:spTree/p:sp/p:txBody/a:p'


class Describe_Run(object):

    def it_provides_access_to_the_font_of_the_run(
            self, r_, _Font_, rPr_, font_):
        run = _Run(r_, None)
        font = run.font
        r_.get_or_add_rPr.assert_called_once_with()
        _Font_.assert_called_once_with(rPr_)
        assert font == font_

    def it_provides_access_to_the_hyperlink_of_the_run(self, run):
        hlink = run.hyperlink
        assert isinstance(hlink, _Hyperlink)

    def it_can_get_the_text_of_the_run(self, run, test_text):
        assert run.text == test_text

    def it_can_change_the_text_of_the_run(self, run):
        run.text = 'new text'
        assert run.text == 'new text'

    # fixtures ---------------------------------------------

    @pytest.fixture
    def _Font_(self, request, font_):
        _Font_ = class_mock(request, 'pptx.text._Font')
        _Font_.return_value = font_
        return _Font_

    @pytest.fixture
    def font_(self, request):
        return instance_mock(request, 'pptx.text._Font')

    @pytest.fixture
    def r(self, r_xml):
        return parse_xml_bytes(r_xml)

    @pytest.fixture
    def rPr_(self, request):
        return instance_mock(request, CT_TextCharacterProperties)

    @pytest.fixture
    def r_(self, request, rPr_):
        r_ = instance_mock(request, CT_RegularTextRun)
        r_.get_or_add_rPr.return_value = rPr_
        return r_

    @pytest.fixture
    def r_xml(self, test_text):
        return ('<a:r %s><a:t>%s</a:t></a:r>' %
                (nsdecls('a'), test_text))

    @pytest.fixture
    def run(self, r):
        return _Run(r, None)

    @pytest.fixture
    def test_text(self):
        return 'test text'
